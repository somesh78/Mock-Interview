import os
import re
import threading
from datetime import datetime

from dotenv import load_dotenv
from docx import Document
from fpdf import FPDF
from groq import Groq
import spacy

import cv2
import mediapipe as mp

import pyttsx3
from speechRecognition import capture_audio_and_analyze
from posturedetection import assess_posture

load_dotenv()

# ─── Initialize TTS ─────────────────────────────────────────────────────────────
tts_engine = pyttsx3.init()
tts_engine.setProperty('rate', 150)      # speaking speed
tts_engine.setProperty('volume', 1.0)    # volume 0–1

def speak_async(text: str):
    """Speak the given text in a background thread (non-blocking)."""
    def _worker():
        tts_engine.say(text)
        tts_engine.runAndWait()            # Blocks this thread only :contentReference[oaicite:0]{index=0}
    threading.Thread(target=_worker, daemon=True).start()  # Non-blocking in main thread :contentReference[oaicite:1]{index=1}


# ─── Groq Client & NLP Setup ────────────────────────────────────────────────────
GROQ_API_KEY = os.getenv('GROQ_API_KEY')
client = Groq(api_key=GROQ_API_KEY)
nlp = spacy.load("en_core_web_sm")


# ─── Resume Parsing & Question Generation ───────────────────────────────────────
def extract_resume_info(resume_path):
    doc = Document(resume_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    return extract_projects_and_skills(text)

def extract_projects_and_skills(text, personal_info=None,
                                skill_keywords=None, project_keywords=None):
    if skill_keywords is None:
        skill_keywords = ["machine learning", "deep learning", "python",
                          "data analysis", "AI", "neural networks"]
    if project_keywords is None:
        project_keywords = ["project", "developed",
                            "implemented", "designed", "built"]
    personal_info = personal_info or []
    for info in personal_info:
        text = re.sub(re.escape(info), "", text, flags=re.IGNORECASE)

    skills, projects = [], []
    for line in text.splitlines():
        if any(sk.lower() in line.lower() for sk in skill_keywords):
            skills.append(line.strip())
        if any(pk.lower() in line.lower() for pk in project_keywords):
            projects.append(line.strip())
    return skills, projects

def generate_custom_questions(projects, skills):
    qs = []
    for proj in projects:
        qs.append(f"Can you tell me more about your experience working on '{proj}'? What was your role?")
    for sk in skills:
        qs.append(f"How have you applied '{sk}' in your previous projects? Give an example.")
    return qs


# ─── Feedback & PDF Generation ─────────────────────────────────────────────────
def get_feedback(user_answer, correct_answer):
    prompt = (
        "The user provided the following answer:\n"
        f"{user_answer}\n\n"
        "Compare it with this correct answer:\n"
        f"{correct_answer}\n\n"
        "Provide constructive feedback."
    )
    resp = client.chat.completions.create(
        messages=[{"role": "user", "content": prompt}],
        model="llama3-8b-8192",
    )
    return resp.choices[0].message.content.strip()

def get_resume_feedback(resume_text):
    prompt = (
        "Based on this resume content:\n"
        f"{resume_text}\n\n"
        "Provide feedback on structure, clarity, skills, and improvements."
    )
    resp = client.chat.completions.create(
        messages=[{"role": "user", "content": prompt}],
        model="llama3-8b-8192",
    )
    return resp.choices[0].message.content.strip()

def generate_pdf(questions, pdf_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(0, 10, txt="Interview Questions", ln=True, align="C")
    for i, q in enumerate(questions, 1):
        pdf.multi_cell(0, 10, txt=f"Question {i}: {q}")
    pdf.output(pdf_path)


# ─── Posture Recording ─────────────────────────────────────────────────────────
posture_data = {}
def record_posture_data(posture, eye_movement):
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    posture_data[ts] = {"posture": posture, "eye_movement": eye_movement}


# ─── Main Interview Flow ───────────────────────────────────────────────────────
def conduct_interview(resume_path, pdf_path):
    # 1. Welcome
    print("\nWelcome to the Personalized Interview Bot!\n")
    speak_async("Welcome to your personalized mock interview.")  # non-blocking :contentReference[oaicite:2]{index=2}

    # 2. Parse resume
    personal_info = ["Your Name", "your.email@example.com", "Your University"]
    skills, projects = extract_resume_info(resume_path)
    if not (skills or projects):
        print("No skills/projects found; using general questions.")
        speak_async("No specialized experience detected; I will ask general questions.")
        questions = ["Tell me about yourself.", "What are your strengths and weaknesses?"]
    else:
        questions = generate_custom_questions(projects, skills)

    # 3. Initialize video capture for posture
    mp_pose = mp.solutions.pose
    pose = mp_pose.Pose(min_detection_confidence=0.5, min_tracking_confidence=0.5)
    cap = cv2.VideoCapture(0)

    all_q, answers, corrects = [], [], []

    for idx, q in enumerate(questions, 1):
        prompt = f"Question {idx}: {q}"
        print(prompt)
        speak_async(prompt)                              # spoken on background thread :contentReference[oaicite:3]{index=3}
        all_q.append(q)

        # 4. Capture audio answer
        text, tone = capture_audio_and_analyze()          # records & transcribes :contentReference[oaicite:4]{index=4}
        if text is None:
            print("Didn't catch that; please repeat.")
            speak_async("I didn't catch that. Please repeat.")
            text, tone = capture_audio_and_analyze()
        print("Answer:", text)
        if tone:
            print("Tone detected:", tone)

        answers.append(text)

        # 5. Sample posture
        ret, frame = cap.read()
        if ret:
            results = pose.process(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB))
            if results.pose_landmarks:
                pf, em = assess_posture(results.pose_landmarks.landmark)
                record_posture_data(pf, em)
                print("Posture feedback:", pf)
        cv2.imshow("Interview Feed", frame)
        cv2.waitKey(1)                                    # prevents OpenCV freeze :contentReference[oaicite:5]{index=5}

        # 6. Generate model answer
        cprompt = f"Provide a detailed, structured answer for: {q}"
        cresp = client.chat.completions.create(
            messages=[{"role": "user", "content": cprompt}],
            model="llama3-8b-8192",
        )
        ca = cresp.choices[0].message.content.strip()
        corrects.append(ca)

    # Cleanup
    cap.release()
    pose.close()
    cv2.destroyAllWindows()

    # 7. Session end
    print("\nInterview ended. Thank you!\n")
    speak_async("The mock interview has concluded. Thank you for your time.")  # non-blocking :contentReference[oaicite:6]{index=6}

    # 8. Per-question feedback
    for i, (ua, ca) in enumerate(zip(answers, corrects), 1):
        fb = get_feedback(ua, ca)
        print(f"\nFeedback for Q{i}:\n{fb}\n")
        speak_async(f"Feedback for question {i}: {fb}")    # background :contentReference[oaicite:7]{index=7}

    # 9. Posture summary
    print("Posture & eye movement:")
    for ts, d in posture_data.items():
        print(f" - {ts}: posture={d['posture']}, eye_movement={d['eye_movement']}")

    # 10. Resume feedback
    raw = open(resume_path, encoding='utf-8', errors='ignore').read()
    rf = get_resume_feedback(raw)
    print("\nResume Feedback:\n", rf)
    speak_async("Here is feedback on your resume.")      # non-blocking :contentReference[oaicite:8]{index=8}

    # 11. Save PDF
    generate_pdf(all_q, pdf_path)
    print(f"\nQuestions saved to '{pdf_path}'.")
    speak_async("I have saved all questions into a PDF for your review.")  # :contentReference[oaicite:9]{index=9}

if __name__ == "__main__":
    resume_path = input("Enter path to your resume (.docx): ").strip()
    conduct_interview(resume_path, "interview_questions.pdf")
    print("Mock interview completed. Check the PDF for your questions.")